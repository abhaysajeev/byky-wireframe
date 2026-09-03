"""Build BYKY wireframe seed data from the client Excel files.

Explicit alias map (no fuzzy matching) so every reconciliation decision is auditable.
"""
import openpyxl, json, re, os
from collections import Counter, OrderedDict

DOCS = "/home/silpc-068/Desktop/2026/byky-main/byky Docs"

# Station name variants: vehicle-sheet spelling -> locations-sheet spelling.
ALIASES = {
    "MAMZAR PARK": "AL MAMZAR PARK",
    "BARSHA POND PARK": "AL BARSHA POND PARK",
    "ZABEEL PARK 1": "AL ZABEEL PARK GATE-1",
    "AL ALAM PARK": "ALALAM PARK",
    "AL WADI PARK": "AL WADI",
    "RAK CORNICHE - 1": "RAK COR -1",
    "MINA AL ARAB RAK": "MINA AL ARAB",
    "CREEK PARK 1": "CREEK PARK GATE 1",
    "CREEK PARK 4": "CREEK PARK GATE 4",
    "ABUDHABI CORNICHE - 1": "ABUDHBI CORNICHE -1",
    "ABUDHABI CORNICHE - 2": "ABUDHBI CORNICHE -2",
    "ABUDHABI CORNICHE - 3": "ABUDHBI CORNICHE -3",
    "MAJAZ 1": "MAJAZ - 1",
    "MAZYAD WALKWAY": "MAZYAD WALKWAY",
    "ASHARIJ WALK WAY": "ASHARIJ WALK WAY",
}
# Stations present only in the vehicle sheet, with the emirate assigned from the name.
EXTRA_STATIONS = {
    "AJMAN CORNICHE": "Ajman",
    "HILTON KUWAIT RESORT": "Kuwait",
    "DUBAI": "Dubai",
}
EMIRATE_LABEL = {
    "Dubai": "Dubai", "Abudhabi": "Abu Dhabi", "Sharjah": "Sharjah",
    "Ajman": "Ajman", "Fujairah": "Fujairah", "RAK": "Ras Al Khaimah",
    "Alain": "Al Ain", "Kuwait": "Kuwait",
}

KEEP_UPPER = {"RAK", "PRO", "GM", "UAE"}

def title(s):
    s = " ".join(str(s).split())
    return " ".join(w if w.upper() in KEEP_UPPER else w.capitalize() for w in s.split())

# Spelling fixes for names the client will actually read on screen.
SPELLING = [
    ("Abudhbi", "Abu Dhabi"), ("Abudhabi", "Abu Dhabi"),
    ("Sooter", "Scooter"), ("Electrice", "Electric"),
    ("Kepper", "Keeper"), ("Ifun", "iFun"),
    (r"\bCor\b", "Corniche"), ("Electricbike", "Electric Bike"),
    ("Xploier", "Xplorer"), ("Sun Beem", "Sun Beam"), ("Baraha Baby", "Barsha Baby"),
]

# Vehicle types that are the same model spelled differently across rows.
VTYPE_MERGE = {
    "Two Wheels": "Two Wheel", "Single Bike": "Electric Bike Single",
    "Double Bike": "Electric Bike Double", "Rent Bike": "Rent Bike",
}

def pretty(name):
    out = title(name)
    for bad, good in SPELLING:
        out = re.sub(bad, good, out, flags=re.I)
    out = re.sub(r"\s*-\s*(\d)", r" \1", out)     # "Corniche -2" -> "Corniche 2"
    return " ".join(out.split())

def canon(name):
    n = " ".join(str(name).upper().split())
    return ALIASES.get(n, n)

# ---- Locations -----------------------------------------------------------
wb = openpyxl.load_workbook(f"{DOCS}/Location and Employee Data.xlsx", read_only=True, data_only=True)
loc_rows = [r for r in wb["Locations"].iter_rows(min_row=2, values_only=True) if r[1]]
stations = OrderedDict()
for i, r in enumerate(loc_rows, 1):
    key = " ".join(str(r[1]).upper().split())
    stations[key] = {
        "code": f"ST{i:03d}", "name": pretty(r[1]),
        "emirate": EMIRATE_LABEL.get(str(r[2]).strip(), str(r[2]).strip()),
        "source": "locations",
    }

# ---- Vehicles ------------------------------------------------------------
wbv = openpyxl.load_workbook(f"{DOCS}/VehicleDetails.xlsx", read_only=True, data_only=True)
vrows = [r for r in wbv["Sheet1"].iter_rows(min_row=2, values_only=True) if r[0]]

for r in vrows:
    key = canon(r[0])
    if key not in stations:
        em = EXTRA_STATIONS.get(key)
        if em is None:
            continue
        stations[key] = {"code": f"ST{len(stations)+1:03d}", "name": pretty(key),
                         "emirate": em, "source": "vehicles"}

vehicles = []
for r in vrows:
    key = canon(r[0])
    st = stations.get(key)
    vehicles.append({
        "number": str(r[3]).strip().upper(),
        "barcode": str(r[4]).strip(),
        "category": pretty(r[1]),
        "vtype": VTYPE_MERGE.get(pretty(r[2]), pretty(r[2])),
        "station": st["name"] if st else pretty(r[0]),
        "station_code": st["code"] if st else None,
    })

counts = Counter(v["station_code"] for v in vehicles)
for s in stations.values():
    s["vehicle_count"] = counts.get(s["code"], 0)

# ---- Staff ---------------------------------------------------------------
srows = [r for r in wb["Staff"].iter_rows(min_row=2, values_only=True) if r[1]]
PROF = {"LABOUR": "Labour", "G . M": "General Manager", "AST.ACCOUNTANT": "Assistant Accountant",
        "STORE KEPPER": "Store Keeper", "PRO": "PRO"}
employees = []
for r in srows:
    p = str(r[3]).strip()
    employees.append({
        "emp_no": str(r[1]).strip(),
        "name": pretty(r[2]),
        "profession": PROF.get(p.upper(), PROF.get(p, pretty(p))),
    })

out = {
    "stations": list(stations.values()),
    "employees": employees,
    "vehicles": vehicles,
    "categories": sorted({v["category"] for v in vehicles}),
    "vehicle_types": sorted({v["vtype"] for v in vehicles}),
    "emirates": sorted({s["emirate"] for s in stations.values()}),
}
dest = "/home/silpc-068/Desktop/2026/byky-main/django-version/full-version/apps/byky_core/seed_data.json"
json.dump(out, open(dest, "w"), indent=1)
print("stations  ", len(out["stations"]))
print("employees ", len(out["employees"]))
print("vehicles  ", len(out["vehicles"]))
print("emirates  ", out["emirates"])
print("categories", out["categories"])
print("unmapped vehicles:", sum(1 for v in vehicles if v["station_code"] is None))
print("stations with 0 vehicles:", [s["name"] for s in out["stations"] if s["vehicle_count"] == 0])
